"""قراءة ملفات استيراد فاتورة الشراء: Excel (xlsx) / CSV / JSON — مطابق لـ handlePurchaseImportFile الأصلي.

JSON المتوقع:  {"supplier": "...", "supplierInvoiceNumber": "...", "notes": "...",
                "items": [{"name","code","qty","cost","price","priceWholesale","priceDistribution"}]}
Excel/CSV: الصف الأول ترويسة (اسم الصنف، الكود، الكمية، السعر) وتُكتشف الأعمدة تلقائياً.
"""
from __future__ import annotations

import csv
import io
import json
from dataclasses import dataclass, field

from .search import RawRow, parse_table_rows


class ImportError_(ValueError):
    """خطأ قراءة ملف — الرسالة عربية وجاهزة للعرض."""


@dataclass
class ImportPayload:
    items: list[RawRow] = field(default_factory=list)
    supplier: str = ""
    supplier_invoice_number: str = ""
    notes: str = ""


def _rows_from_xlsx(data: bytes) -> list[list]:
    try:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        return [["" if c is None else c for c in row] for row in ws.iter_rows(values_only=True)]
    except ImportError as e:
        raise ImportError_("ثبّت openpyxl لقراءة ملفات Excel") from e
    except Exception as e:  # noqa: BLE001
        raise ImportError_("تعذّر قراءة ملف الإكسل — تأكد من الملف وحاول مرة ثانية") from e


def _rows_from_csv(data: bytes) -> list[list]:
    for enc in ("utf-8-sig", "cp1256", "latin-1"):  # cp1256 = ترميز العربية القديم لملفات Excel/CSV
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    dialect = csv.Sniffer().sniff(text[:2048], delimiters=",;\t") if text.strip() else csv.excel
    return [row for row in csv.reader(io.StringIO(text), dialect)]


def parse_import_file(filename: str, data: bytes) -> ImportPayload:
    name = (filename or "").lower()
    if name.endswith(".xls"):
        raise ImportError_("صيغة xls القديمة غير مدعومة — احفظ الملف كـ xlsx أو csv ثم أعد المحاولة")
    if name.endswith((".xlsx", ".csv")):
        rows = _rows_from_xlsx(data) if name.endswith(".xlsx") else _rows_from_csv(data)
        items = parse_table_rows(rows)
        if not items:
            raise ImportError_("ما لقيت أصناف بالملف — تأكد أن الأعمدة: اسم الصنف، الكود، الكمية، السعر")
        return ImportPayload(items=items)
    try:
        obj = json.loads(data.decode("utf-8-sig"))
    except (UnicodeDecodeError, ValueError) as e:
        raise ImportError_("الملف ليس JSON صحيحاً — تأكد من الملف وحاول مرة ثانية") from e
    if not isinstance(obj, dict) or not isinstance(obj.get("items"), list) or not obj["items"]:
        raise ImportError_("صيغة الملف غير متوقعة — لازم يحتوي على items كمصفوفة أصناف")
    s = lambda v: "" if v is None else str(v)  # noqa: E731
    rows = [RawRow(s(r.get("name")).strip(), s(r.get("code")).strip(), s(r.get("qty")), s(r.get("cost")),
                   s(r.get("price")), s(r.get("priceWholesale")), s(r.get("priceDistribution")))
            for r in obj["items"] if isinstance(r, dict)]
    return ImportPayload(rows, s(obj.get("supplier")), s(obj.get("supplierInvoiceNumber")), s(obj.get("notes")))
