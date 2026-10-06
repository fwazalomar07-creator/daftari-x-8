"""تصدير فاتورة كاملة إلى Excel: الشعار + اسم الشركة والبيانات + البنود مع صورة كل صنف + المجاميع (+ الباركود).

الشكل مطابق للفاتورة المطبوعة: ترويسة كريمية وخطوط ذهبية، أعمدة (من اليمين): الصنف، الكود، الكمية، السعر، الإجمالي، الصورة.
الأرقام تُكتب أرقاماً حقيقية، وإجمالي كل بند معادلة (=الكمية×السعر) فيمكن تعديلها داخل Excel.
"""
from __future__ import annotations

import io
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from PIL import Image, ImageOps

from ..core.models import Invoice
from .render import ASSETS, arabic_date, resolve_barcode

FONT = "Arial"
GOLD, CREAM, INK, RED = "B8862E", "F6EEDA", "1B2E2A", "B0402F"
MONEY = '"$"#,##0.00'
IMG_PX = 72                       # أبعاد صورة الصنف داخل الخلية (بكسل)
ROW_PT = 60                       # ارتفاع صف البند (نقطة) ≈ 80 بكسل


def _img_for_cell(data: bytes, px: int) -> XLImage | None:
    try:
        im = Image.open(io.BytesIO(data)).convert("RGB")
        im = ImageOps.contain(im, (px, px), Image.LANCZOS)
        canvas = Image.new("RGB", (px, px), (255, 255, 255))
        canvas.paste(im, ((px - im.width) // 2, (px - im.height) // 2))
        buf = io.BytesIO()
        canvas.save(buf, "PNG")
        buf.seek(0)
        x = XLImage(buf)
        x.width = x.height = px
        return x
    except Exception:  # noqa: BLE001 — صورة تالفة لا توقف التصدير
        return None


def _logo_image(width_px: int = 640) -> XLImage | None:
    path = next(iter(ASSETS.glob("logo.*")), None)
    if not path:
        return None
    try:
        im = Image.open(path).convert("RGB")
        h = int(im.height * width_px / im.width)
        im = im.resize((width_px, h), Image.LANCZOS)
        buf = io.BytesIO()
        im.save(buf, "PNG")
        buf.seek(0)
        x = XLImage(buf)
        x.width, x.height = width_px, h
        return x
    except Exception:  # noqa: BLE001
        return None


def invoice_xlsx(inv: Invoice, settings: dict, images: dict[str, bytes] | None = None,
                 barcode: bytes | None = None) -> bytes:
    images = images or {}
    with_imgs = settings.get("printImages") is not False
    wb = Workbook()
    ws = wb.active
    ws.title = "فاتورة"[:31]
    ws.sheet_view.rightToLeft = True
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToHeight = 0

    # الأعمدة بالترتيب العربي: A=الصنف … F=الصورة (الورقة RTL فيظهر A أقصى اليمين)
    for col, wd in zip("ABCDEF", (38, 16, 10, 13, 15, 14 if with_imgs else 2)):
        ws.column_dimensions[col].width = wd
    last = "F" if with_imgs else "E"
    thin = Side(style="thin", color="E4E0D6")
    gold = Side(style="medium", color=GOLD)

    r = 1
    logo = _logo_image()
    if logo:
        ws.add_image(logo, "A1")
        rows_for_logo = max(4, int(logo.height / 20) + 1)       # 20 بكسل ≈ ارتفاع الصف الافتراضي
        r = rows_for_logo + 1
    # اسم الشركة والبيانات
    def line(text, size=11, bold=False, color=INK):
        nonlocal r
        if not text:
            return
        ws.merge_cells(f"A{r}:{last}{r}")
        c = ws[f"A{r}"]
        c.value = text
        c.font = Font(name=FONT, size=size, bold=bold, color=color)
        c.alignment = Alignment(horizontal="right", vertical="center", readingOrder=2)
        ws.row_dimensions[r].height = size * 1.9
        r += 1
    line(settings.get("businessName"), 16, True)
    line(settings.get("tagline"), 11, False, "5B6C67")
    line(settings.get("address"), 11, False, "5B6C67")
    if settings.get("phone"):
        line(f"هاتف المحل: {settings['phone']}", 11, False, "5B6C67")
    if settings.get("companyNumber"):
        line(f"رقم الشركة: {settings['companyNumber']}", 11, False, "5B6C67")
    r += 1
    line(f"فاتورة رقم: {inv.number}", 12, True)
    line(f"التاريخ: {arabic_date(inv.date)}", 11)
    if inv.customer:
        line(f"العميل: {inv.customer}", 11)
    for c in range(1, 7 if with_imgs else 6):
        ws.cell(row=r, column=c).border = Border(top=gold)
    r += 1

    # ترويسة الجدول
    heads = ["الصنف", "الكود", "الكمية", "السعر", "الإجمالي"] + (["الصورة"] if with_imgs else [])
    for i, h in enumerate(heads, 1):
        c = ws.cell(row=r, column=i, value=h)
        c.font = Font(name=FONT, size=11, bold=True, color=INK)
        c.fill = PatternFill("solid", fgColor=CREAM)
        c.alignment = Alignment(horizontal="right" if i == 1 else "center", vertical="center", readingOrder=2)
        c.border = Border(bottom=thin)
    ws.row_dimensions[r].height = 24
    r += 1
    first = r

    for l in inv.items:
        vals = [l.name, l.code or "—", int(l.qty), float(l.price), f"=C{r}*D{r}"]
        for i, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=i, value=v)
            c.font = Font(name=FONT, size=11, color=INK)
            c.alignment = Alignment(horizontal="right" if i == 1 else "center", vertical="center",
                                    wrap_text=True, readingOrder=2 if i == 1 else 0)
            c.border = Border(bottom=thin)
            if i in (4, 5):
                c.number_format = MONEY
        if with_imgs:
            ws.cell(row=r, column=6).border = Border(bottom=thin)
            ws.row_dimensions[r].height = ROW_PT
            data = images.get(getattr(l, "id", None))
            x = _img_for_cell(data, IMG_PX) if data else None
            if x:
                ws.add_image(x, f"F{r}")
        else:
            ws.row_dimensions[r].height = 22
        r += 1
    lastrow = r - 1
    r += 1

    # المجاميع
    def total(label, value, color=INK, grand=False, formula=None):
        nonlocal r
        a = ws.cell(row=r, column=1, value=label)
        b = ws.cell(row=r, column=2, value=formula or float(value))
        for c in (a, b):
            c.font = Font(name=FONT, size=12 if grand else 11, bold=grand, color=color)
            if grand:
                c.border = Border(top=gold)
        a.alignment = Alignment(horizontal="right", readingOrder=2)
        b.alignment = Alignment(horizontal="left")
        b.number_format = MONEY
        r += 1
    if inv.discount:
        total("المجموع الفرعي", inv.subtotal)
        total("الخصم", -inv.discount, RED)
    total("الإجمالي الكلي", inv.total, INK, True, f"=SUM(E{first}:E{lastrow})-{float(inv.discount or 0)}" if inv.items else None)
    total("المدفوع الآن", inv.paid)
    if inv.customer and inv.customer_prior_debt:
        total("دين سابق على العميل", inv.customer_prior_debt, RED)
    if inv.remaining > 0:
        total("المتبقي من هذه الفاتورة", inv.remaining, RED)
    if inv.customer and inv.customer_total_debt_after > 0:
        total("إجمالي الدين على العميل الآن", inv.customer_total_debt_after, RED, True)
    if inv.notes:
        r += 1
        line(f"ملاحظات: {inv.notes}")

    bc = resolve_barcode(inv, settings, override=barcode)
    if bc:
        try:
            im = Image.open(io.BytesIO(bc)).convert("RGB")
            im = ImageOps.contain(im, (360, 140), Image.LANCZOS)
            buf = io.BytesIO()
            im.save(buf, "PNG")
            buf.seek(0)
            x = XLImage(buf)
            x.width, x.height = im.size
            r += 1
            ws.add_image(x, f"A{r}")
        except Exception:  # noqa: BLE001
            pass

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
