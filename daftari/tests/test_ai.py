"""اختبارات المساعد: بيئة .env، عميل Gemini (نقل وهمي)، الأدوات فوق Ledger، وحلقة الوكيل."""
import asyncio, json, os, tempfile, unittest
from decimal import Decimal as D
from pathlib import Path

from daftari.ai.agent import Assistant
from daftari.ai.gemini import GeminiClient, GeminiConfig, GeminiError
from daftari.ai.tools import ToolBox, jsonable
from daftari.core.env import load_env, parse_env_text
from daftari.data.repository import Repository
from daftari.data.store import MemoryStore, StoreError
from daftari.ledger import Ledger


def reply_json(parts, status=200):
    body = {"candidates": [{"content": {"role": "model", "parts": parts}, "finishReason": "STOP"}]}
    return status, json.dumps(body).encode()


class FakeTransport:
    """يعيد ردوداً مبرمجة ويسجّل الطلبات."""
    def __init__(self, *responses):
        self.responses, self.calls = list(responses), []

    async def __call__(self, url, headers, body, timeout):
        self.calls.append((url, headers, json.loads(body)))
        return self.responses.pop(0)


class EnvTests(unittest.TestCase):
    def test_parse(self):
        d = parse_env_text('# c\nA=1\nexport B="two words"\nC=x # tail\n\nBAD\nD=\n')
        self.assertEqual(d, {"A": "1", "B": "two words", "C": "x", "D": ""})

    def test_load_does_not_override_and_skips_empty(self):
        with tempfile.TemporaryDirectory() as t:
            p = Path(t) / ".env"
            p.write_text("ZZ_TEST_A=file\nZZ_TEST_B=\nZZ_TEST_C=c\n", encoding="utf-8")
            os.environ["ZZ_TEST_A"] = "real"; os.environ.pop("ZZ_TEST_B", None); os.environ.pop("ZZ_TEST_C", None)
            load_env([p])
            self.assertEqual(os.environ["ZZ_TEST_A"], "real")
            self.assertNotIn("ZZ_TEST_B", os.environ)
            self.assertEqual(os.environ["ZZ_TEST_C"], "c")
            for k in ("ZZ_TEST_A", "ZZ_TEST_C"):
                os.environ.pop(k, None)


class GeminiClientTests(unittest.IsolatedAsyncioTestCase):
    def cfg(self, **kw):
        return GeminiConfig(api_key="K", model="m1", **kw)

    async def test_request_shape_and_parse(self):
        tr = FakeTransport(reply_json([{"text": "مرحبا"}]))
        c = GeminiClient(self.cfg(), tr)
        r = await c.generate([{"role": "user", "parts": [{"text": "hi"}]}], "sys", [{"name": "t", "description": "d"}])
        self.assertEqual(r.text, "مرحبا")
        url, headers, body = tr.calls[0]
        self.assertTrue(url.endswith("/models/m1:generateContent"))
        self.assertEqual(headers["x-goog-api-key"], "K")
        self.assertEqual(body["systemInstruction"]["parts"][0]["text"], "sys")
        self.assertEqual(body["generationConfig"]["thinkingConfig"]["thinkingLevel"], "high")
        self.assertEqual(body["tools"][0], {"googleSearch": {}})
        self.assertEqual(body["tools"][1]["functionDeclarations"][0]["name"], "t")

    async def test_transcribe_sends_audio_without_google_tool(self):
        tr = FakeTransport(reply_json([{"text": " مرحبا بالجميع "}]))
        c = GeminiClient(self.cfg(), tr)
        text = await c.transcribe_audio(b"RIFF" + b"\x00" * 12, "audio/wav")
        self.assertEqual(text, "مرحبا بالجميع")
        body = tr.calls[0][2]
        self.assertNotIn("tools", body)
        self.assertEqual(body["contents"][0]["parts"][0]["inlineData"]["mimeType"], "audio/wav")

    async def test_function_call_parsed_and_thought_signature_kept(self):
        part = {"functionCall": {"name": "x", "args": {"a": 1}}, "thoughtSignature": "SIG"}
        r = await GeminiClient(self.cfg(), FakeTransport(reply_json([part]))).generate([])
        self.assertEqual(r.calls, [{"name": "x", "args": {"a": 1}}])
        self.assertEqual(r.content["parts"][0]["thoughtSignature"], "SIG")

    async def test_retry_on_429_then_ok(self):
        tr = FakeTransport((429, b"{}"), reply_json([{"text": "ok"}]))
        r = await GeminiClient(self.cfg(), tr, backoff=0).generate([])
        self.assertEqual((r.text, len(tr.calls)), ("ok", 2))

    async def test_thinking_rejected_falls_back(self):
        err = (400, json.dumps({"error": {"message": "Unknown field thinkingLevel (thinking)"}}).encode())
        tr = FakeTransport(err, reply_json([{"text": "ok"}]))
        r = await GeminiClient(self.cfg(), tr).generate([])
        self.assertEqual(r.text, "ok")
        self.assertNotIn("thinkingConfig", tr.calls[1][2]["generationConfig"])

    async def test_errors_are_explained_in_arabic(self):
        for status, word in ((403, "مرفوض"), (404, "GEMINI_MODEL"), (429, "حصة")):
            with self.assertRaises(GeminiError) as cm:
                await GeminiClient(self.cfg(), FakeTransport((status, b"{}")), retries=0).generate([])
            self.assertIn(word, str(cm.exception))
        with self.assertRaises(GeminiError):
            await GeminiClient(GeminiConfig(), FakeTransport()).generate([])  # بلا مفتاح


class ToolBoxTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.home = Path(self.tmp.name)
        self.store = MemoryStore()
        self.led = Ledger(Repository(self.store, cache_dir=self.home / "c"))
        await self.led.load()
        self.item = await self.led.add_item("فلتر زيت", "F1", 5, 9, 10)
        await self.led.add_customer("أبو أحمد", "0999")
        self.asked: list[str] = []
        self.answer = True

        async def confirm(msg):
            self.asked.append(msg)
            return self.answer
        self.tb = ToolBox(self.led, self.store, self.home, confirm)

    async def asyncTearDown(self):
        self.tmp.cleanup()

    async def test_declarations_are_valid_for_gemini(self):
        decls = self.tb.declarations()
        names = [d["name"] for d in decls]
        self.assertEqual(len(names), len(set(names)))
        for d in decls:
            self.assertTrue(d["description"])
            def check(node):
                self.assertIn(node["type"], {"STRING", "NUMBER", "INTEGER", "BOOLEAN", "ARRAY", "OBJECT"})
                if node["type"] == "ARRAY":
                    self.assertIn("items", node)               # Gemini يرفض ARRAY بلا items
                    check(node["items"])
                for sub in node.get("properties", {}).values():
                    check(sub)
            for p in d.get("parameters", {}).get("properties", {}).values():
                check(p)
        json.dumps(decls)

    async def test_export_dataset_inventory_to_real_xlsx(self):
        from openpyxl import load_workbook
        out = await self.tb.execute("export_dataset", {"dataset": "inventory", "filename": "مخزون"})
        self.assertTrue(out["ok"], out)
        ws = load_workbook(out["result"]["path"]).worksheets[0]
        self.assertTrue(ws.sheet_view.rightToLeft)
        self.assertEqual([c.value for c in ws[1]][:3], ["الصنف", "الكود", "الكمية"])
        self.assertEqual(ws["A2"].value, "فلتر زيت")
        self.assertEqual(ws["C2"].value, 10)                       # كمية رقمية حقيقية
        self.assertEqual(ws["E2"].value, 9)                        # سعر المفرق رقم لا نص
        self.assertIn("$", ws["E2"].number_format)
        self.assertEqual(ws["A3"].value, "الإجمالي")               # صف الإجماليات
        self.assertTrue(str(ws["D3"].value).startswith("=SUM("))
        self.assertEqual(len(self.tb.exports), 1)

    async def test_export_excel_from_model_tables_and_errors(self):
        from openpyxl import load_workbook
        out = await self.tb.execute("export_excel", {"filename": "تحليل: أرباح/شهر", "sheets": [
            {"name": "أرباح [2026]", "headers": ["الشهر", "الربح", "النسبة"],
             "rows": [["أكتوبر", "$1,250.5", "12%"], ["نوفمبر", "980", "—"]]},
            {"name": "أرباح [2026]", "headers": ["x"], "rows": [["1"]]}]})
        self.assertTrue(out["ok"], out)
        self.assertNotIn(":", out["result"]["file"])               # اسم ملف آمن لويندوز
        wb = load_workbook(out["result"]["path"])
        self.assertEqual(len(set(wb.sheetnames)), 2)               # أسماء الأوراق فريدة وبلا رموز ممنوعة
        ws = wb.worksheets[0]
        self.assertEqual((ws["B2"].value, ws["C2"].value), (1250.5, 0.12))
        bad = await self.tb.execute("export_excel", {"filename": "x", "sheets": []})
        self.assertFalse(bad["ok"])
        bad2 = await self.tb.execute("export_dataset", {"dataset": "nope"})
        self.assertFalse(bad2["ok"])
        empty = await self.tb.execute("export_dataset", {"dataset": "debtors"})   # لا مدينين
        self.assertFalse(empty["ok"])

    async def test_read_tools(self):
        s = (await self.tb.execute("business_summary", {}))["result"]
        self.assertEqual(s["items"], 1)
        found = (await self.tb.execute("search_items", {"query": "فلتر"}))["result"]
        self.assertEqual(found["items"][0]["code"], "F1")
        json.dumps(await self.tb.execute("list_customers", {}))
        self.assertFalse((await self.tb.execute("get_invoice", {"number": "nope"}))["ok"])

    async def test_write_requires_confirmation_and_makes_backup(self):
        out = await self.tb.execute("add_item", {"name": "بواجي", "cost": 3, "price": 5, "stock": 4, "code": "SP"})
        self.assertTrue(out["ok"]); self.assertEqual(len(self.asked), 1)
        self.assertTrue(any(i.name == "بواجي" for i in self.led.inventory))
        self.assertEqual(len(list((self.home / "backups").glob("*.json"))), 1)

    async def test_rejected_write_changes_nothing(self):
        self.answer = False
        before = len(self.led.inventory)
        out = await self.tb.execute("add_item", {"name": "x", "cost": 1, "price": 2})
        self.assertTrue(out.get("rejected")); self.assertEqual(len(self.led.inventory), before)
        self.assertFalse((self.home / "backups").exists())

    async def test_auto_approve_skips_prompt(self):
        tb = ToolBox(self.led, self.store, self.home, None, lambda: True)
        self.assertTrue((await tb.execute("restock_item", {"code": "F1", "qty": 5}))["ok"])
        self.assertEqual(self.led.inventory[0].stock, 15)

    async def test_price_edit_keeps_unspecified(self):
        await self.tb.execute("edit_item_prices", {"item": "F1", "retail_price": 12})
        it = self.led.inventory[0]
        self.assertEqual((it.cost, it.price, it.price_wholesale), (D(5), D(12), D(9)))

    async def test_payment_and_validation_errors_are_returned_not_raised(self):
        await self.led.add_manual_customer_debt(self.led.customers[0].id, 50)
        out = await self.tb.execute("record_customer_payment", {"customer_name": "أبو أحمد", "amount": 20})
        self.assertTrue(out["ok"]); self.assertEqual(self.led.customers[0].balance, D(30))
        bad = await self.tb.execute("add_expense", {"category": "zzz", "amount": 5})
        self.assertFalse(bad["ok"])
        self.assertFalse((await self.tb.execute("add_item", {"bogus": 1}))["ok"])
        self.assertFalse((await self.tb.execute("no_such_tool", {}))["ok"])

    async def test_google_search_tool_parses_results(self):
        from daftari.ai import websearch as ws
        html = '<html><a href="/url?q=https://example.com/filter">فلتر زيت تويوتا سعر</a></html>'

        def fake_fetch(url, body, headers, timeout):
            if "google.com" in url:
                return 200, html
            return 200, ""

        orig = ws.default_fetch
        ws.default_fetch = fake_fetch
        try:
            out = await self.tb.execute("google_search", {"query": "فلتر زيت تويوتا"})
        finally:
            ws.default_fetch = orig
        self.assertTrue(out["ok"], out)
        self.assertEqual(out["result"]["engine"], "google")
        self.assertEqual(out["result"]["results"][0]["url"], "https://example.com/filter")
        self.assertIn("google_search", [d["name"] for d in self.tb.declarations()])

    async def test_raw_data_and_project_files_are_read_only_and_safe(self):
        keys = (await self.tb.execute("list_data_keys", {}))["result"]["keys"]
        self.assertIn("inventory", keys)
        raw = (await self.tb.execute("read_data_key", {"key": "inventory"}))["result"]
        self.assertIn("فلتر", raw["json"])
        self.assertTrue((await self.tb.execute("read_project_file", {"path": "ledger.py"}))["ok"])
        for bad in ("../../etc/passwd", ".env", "assets/logo.jpeg"):
            self.assertFalse((await self.tb.execute("read_project_file", {"path": bad}))["ok"], bad)

    async def test_save_proposal_writes_outside_source(self):
        out = await self.tb.execute("save_proposal", {"title": "تقرير الركود", "content": "```py\nprint(1)\n```"})
        self.assertTrue(Path(out["result"]["saved_to"]).read_text(encoding="utf-8").startswith("# تقرير"))
        self.assertTrue(str(out["result"]["saved_to"]).startswith(str(self.home / "proposals")))

    def test_jsonable(self):
        self.assertEqual(jsonable({"a": D("1.50"), "b": [D(2)]}), {"a": 1.5, "b": [2]})


class AgentTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.home = Path(self.tmp.name)
        self.led = Ledger(Repository(MemoryStore(), cache_dir=self.home / "c"))
        await self.led.load()
        await self.led.add_item("زيت", "O1", 20, 30, 7)
        self.tb = ToolBox(self.led, None, self.home, None, lambda: True)

    async def asyncTearDown(self):
        self.tmp.cleanup()

    async def test_tool_loop_then_final_answer(self):
        tr = FakeTransport(
            reply_json([{"functionCall": {"name": "search_items", "args": {"query": "زيت"}}, "thoughtSignature": "S"}]),
            reply_json([{"text": "عندك صنف واحد: زيت"}]))
        events = []
        a = Assistant(GeminiClient(GeminiConfig(api_key="K"), tr), self.tb, "المتجر")
        out = await a.ask("ما أصناف الزيت؟", lambda k, d: events.append(k))
        self.assertEqual(out, "عندك صنف واحد: زيت")
        self.assertEqual(events, ["tool_start", "tool_done"])
        second_request = tr.calls[1][2]["contents"]
        self.assertEqual(second_request[1]["parts"][0]["thoughtSignature"], "S")           # التوقيع أُعيد
        resp = second_request[2]["parts"][0]["functionResponse"]
        self.assertEqual(resp["name"], "search_items"); self.assertTrue(resp["response"]["ok"])

    async def test_failure_leaves_clean_history(self):
        a = Assistant(GeminiClient(GeminiConfig(api_key="K"), FakeTransport((403, b"{}")), retries=0), self.tb)
        with self.assertRaises(GeminiError):
            await a.ask("hi")
        self.assertEqual(a.history, [])

    async def test_step_limit(self):
        call = reply_json([{"functionCall": {"name": "business_summary", "args": {}}}])
        a = Assistant(GeminiClient(GeminiConfig(api_key="K"), FakeTransport(*[call] * 20)), self.tb)
        a.MAX_STEPS = 3
        with self.assertRaises(GeminiError):
            await a.ask("loop")
        self.assertEqual(a.history, [])

    async def test_history_trim_keeps_valid_start(self):
        a = Assistant(GeminiClient(GeminiConfig(api_key="K"), FakeTransport(*[reply_json([{"text": "r"}])] * 30)), self.tb)
        a.MAX_HISTORY = 6
        for i in range(10):
            await a.ask(f"q{i}")
        self.assertLessEqual(len(a.history), 7)
        self.assertEqual(a.history[0]["role"], "user"); self.assertIn("text", a.history[0]["parts"][0])

    async def test_ask_with_image_sends_inline_data(self):
        tr = FakeTransport(reply_json([{"text": "هذه صورة فلتر."}]))
        a = Assistant(GeminiClient(GeminiConfig(api_key="K"), tr), self.tb, "المتجر")
        jpeg = b"\xff\xd8\xff" + b"\x00" * 40
        out = await a.ask("ما هذه القطعة؟", images=[("image/jpeg", jpeg)])
        self.assertEqual(out, "هذه صورة فلتر.")
        parts = tr.calls[0][2]["contents"][0]["parts"]
        self.assertEqual(parts[0]["inlineData"]["mimeType"], "image/jpeg")
        self.assertEqual(parts[1]["text"], "ما هذه القطعة؟")


class RepoDiagnosticsTests(unittest.IsolatedAsyncioTestCase):
    async def test_load_error_and_missing_are_reported(self):
        with tempfile.TemporaryDirectory() as t:
            st = MemoryStore()
            led = Ledger(Repository(st, cache_dir=Path(t)))
            await led.load()
            self.assertIn("inventory", led.health_notice())       # لا سجل inventory
            st.fail = True
            await led.reload(force=True)
            self.assertIn("RLS", led.health_notice())             # فشل قراءة
            self.assertIsNone(Ledger(Repository(None)).health_notice())

    async def test_reload_skips_when_unsynced_changes(self):
        with tempfile.TemporaryDirectory() as t:
            led = Ledger(Repository(MemoryStore(), cache_dir=Path(t)))
            await led.load()
            led.repo.dirty.add("inventory")
            self.assertFalse(await led.reload())
            self.assertTrue(await led.reload(force=True))


if __name__ == "__main__":
    unittest.main()
