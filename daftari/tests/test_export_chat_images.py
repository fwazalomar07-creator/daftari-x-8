"""اختبارات: تحصين تصدير Excel، حفظ الملفات، طباعة/تصدير المحادثة (PDF + Excel)، وإرسال المساعد لصور الأصناف."""
import asyncio
import io
import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image

from daftari.tests import test_ui_smoke as S          # يثبّت flet الوهمية (استيراد الوحدة لا الصنف حتى لا تتكرر اختباراته)
from daftari.tests.test_ai import FakeTransport, reply_json
from daftari.ai import chat_export as ce
from daftari.ai import excel
from daftari.ai.agent import Assistant
from daftari.ai.gemini import GeminiClient, GeminiConfig
from daftari.ai.llm import OpenAIClient, ProviderSpec
from daftari.ai.tools import ToolBox
from daftari.data.images import ImageStore
from daftari.ui.tabs import assistant as assistant_tab


def jpeg(color=(200, 60, 40), size=(120, 90)) -> bytes:
    b = io.BytesIO()
    Image.new("RGB", size, color).save(b, "JPEG")
    return b.getvalue()


def sheet_values(data: bytes, name=None):
    ws = load_workbook(io.BytesIO(data))[name] if name else load_workbook(io.BytesIO(data)).active
    return [[c.value for c in row] for row in ws.iter_rows()]


def last_toast(page) -> str:
    return page.dialogs[-1].args[0].args[0]


# ===================================================================================== تحصين ملف Excel
class ExcelHardening(unittest.TestCase):
    def build(self, rows, headers=("x",), **kw):
        return excel.build_workbook([excel.Sheet(kw.pop("name", "a"), list(headers), rows, **kw)], "العمر")

    def test_control_characters_do_not_crash_and_are_removed(self):
        vals = sheet_values(self.build([["abc\x07def\x0b"], ["مرحبا\x00"]]))
        self.assertEqual(vals[1][0], "abcdef")
        self.assertEqual(vals[2][0], "مرحبا")

    def test_text_starting_with_equals_stays_text_not_formula(self):
        wb = load_workbook(io.BytesIO(self.build([["=1+1"], ['=HYPERLINK("http://x")']], headers=("=bad",))))
        ws = wb.active
        self.assertEqual(ws["A1"].data_type, "s")
        self.assertEqual(ws["A2"].data_type, "s")
        self.assertEqual(ws["A2"].value, "=1+1")

    def test_totals_still_use_real_sum_formulas(self):
        wb = load_workbook(io.BytesIO(self.build([[1.5], [2.5]], money_cols={0}, totals=True)))
        self.assertEqual(wb.active["A4"].value, "=SUM(A2:A3)")
        self.assertEqual(wb.active["A4"].data_type, "f")

    def test_long_text_is_cut_to_excel_cell_limit(self):
        vals = sheet_values(self.build([["x" * 40000]]))
        self.assertEqual(len(vals[1][0]), excel.MAX_CELL_CHARS)

    def test_nan_and_infinity_become_empty_cells(self):
        vals = sheet_values(self.build([[float("nan")], [float("inf")], [1.0]]))
        self.assertEqual([r[0] for r in vals[1:]], [None, None, 1])

    def test_reserved_sheet_name_is_renamed(self):
        wb = load_workbook(io.BytesIO(self.build([["1"]], name="History")))
        self.assertNotEqual(wb.sheetnames[0].lower(), "history")

    def test_normalize_sheets_accepts_every_shape_models_send(self):
        want = [["الصنف", "الكمية"], ["فلتر", 3]]
        shapes = {
            "list of sheets with dict rows": [{"name": "t", "headers": ["الصنف", "الكمية"], "rows": [{"الصنف": "فلتر", "الكمية": 3}]}],
            "dict rows, no headers": [{"name": "t", "rows": [{"الصنف": "فلتر", "الكمية": 3}]}],
            "json string": json.dumps([{"name": "t", "headers": ["الصنف", "الكمية"], "rows": [["فلتر", 3]]}], ensure_ascii=False),
            "single sheet dict": {"name": "t", "headers": ["الصنف", "الكمية"], "rows": [["فلتر", 3]]},
            "wrapped in sheets key": {"sheets": [{"title": "t", "columns": ["الصنف", "الكمية"], "data": [["فلتر", 3]]}]},
            "markdown table text": "| الصنف | الكمية |\n|--|--|\n| فلتر | 3 |\n",
        }
        for label, raw in shapes.items():
            sheets = excel.normalize_sheets(raw)
            self.assertEqual(len(sheets), 1, label)
            self.assertEqual(sheet_values(excel.build_workbook(sheets)), want, label)

    def test_normalize_sheets_rejects_garbage_with_clear_error(self):
        with self.assertRaises(ValueError):
            excel.normalize_sheets("هذا ليس جدولاً")
        self.assertEqual(excel.normalize_sheets([]), [])
        self.assertEqual(excel.normalize_sheets([{"name": "فارغ"}]), [])

    def test_nested_cells_are_serialised_not_crashing(self):
        sheets = excel.normalize_sheets([{"name": "t", "headers": ["a"], "rows": [[{"k": "v"}], [[1, 2]]]}])
        vals = sheet_values(excel.build_workbook(sheets))
        self.assertEqual(vals[1][0], '{"k": "v"}')
        self.assertEqual(vals[2][0], "[1, 2]")


class ExportToolRobustness(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        from daftari.data.repository import Repository
        from daftari.data.store import MemoryStore
        from daftari.ledger import Ledger
        self.tmp = tempfile.TemporaryDirectory()
        self.home = Path(self.tmp.name)
        self.led = Ledger(Repository(MemoryStore(), cache_dir=self.home / "c"))
        await self.led.load()
        self.tb = ToolBox(self.led, None, self.home)

    async def asyncTearDown(self):
        self.tmp.cleanup()

    async def test_export_excel_with_dict_rows_writes_values_not_keys(self):
        out = await self.tb.execute("export_excel", {"filename": "t", "sheets": [
            {"name": "ورقة", "rows": [{"الصنف": "فلتر", "الكمية": 3}, {"الصنف": "زيت", "الكمية": 5}]}]})
        self.assertTrue(out["ok"], out)
        vals = sheet_values(Path(out["result"]["path"]).read_bytes())
        self.assertEqual(vals, [["الصنف", "الكمية"], ["فلتر", 3], ["زيت", 5]])

    async def test_export_excel_with_json_string_sheets(self):
        out = await self.tb.execute("export_excel", {"filename": "t", "sheets": json.dumps(
            [{"name": "a", "headers": ["h"], "rows": [["1"]]}])})
        self.assertTrue(out["ok"], out)

    async def test_export_excel_bad_input_returns_readable_error(self):
        out = await self.tb.execute("export_excel", {"filename": "t", "sheets": "نص عشوائي"})
        self.assertFalse(out["ok"])
        self.assertIn("sheets", out["error"])

    async def test_same_filename_exports_never_overwrite(self):
        paths = set()
        for _ in range(3):
            out = await self.tb.execute("export_excel", {"filename": "same", "sheets": [{"name": "a", "headers": ["h"], "rows": [["1"]]}]})
            self.assertTrue(out["ok"], out)
            paths.add(out["result"]["path"])
        self.assertEqual(len(paths), 3)
        self.assertTrue(all(Path(p).exists() for p in paths))


class LiteProvidersKeepExportTools(unittest.IsolatedAsyncioTestCase):
    async def test_groq_lite_tool_set_includes_export_and_images(self):
        from daftari.data.repository import Repository
        from daftari.data.store import MemoryStore
        from daftari.ledger import Ledger
        with tempfile.TemporaryDirectory() as t:
            led = Ledger(Repository(MemoryStore(), cache_dir=Path(t) / "c"))
            await led.load()
            tb = ToolBox(led, None, Path(t))
            cl = OpenAIClient(ProviderSpec("groq", "openai", "k", "m", "https://api.groq.com/openai/v1"))
            self.assertTrue(cl.lite)
            body = cl.build_body([{"role": "user", "parts": [{"text": "x"}]}], None, tb.declarations())
            names = {t_["function"]["name"] for t_ in body["tools"]}
            self.assertTrue({"export_dataset", "export_excel", "show_item_images"} <= names, names)
            ex = next(t_ for t_ in body["tools"] if t_["function"]["name"] == "export_excel")["function"]["parameters"]
            self.assertEqual(ex["properties"]["sheets"]["items"]["properties"]["rows"]["items"]["items"]["type"], "string")


# ===================================================================================== تصدير المحادثة
class ChatExportTests(unittest.TestCase):
    MD = ("## تحليل\nالقطعة **فلتر** (TY-100).\n\n- موجود بالمخزون\n1. تحقق\n\n"
          "| الصنف | الكمية | السعر |\n|--|--:|--:|\n| فلتر | 6 | $20 |\n| زيت | 30 | $3 |\n\nختام ✅")

    def entries(self):
        msgs = [("user_image", jpeg()), ("user", "📷 صورة مرفقة"), ("assistant", self.MD), ("user", "أرني الصور"),
                ("pictures", [{"id": "a", "name": "فلتر"}]), ("file", "/x/y/مخزون.xlsx"), ("doc", "invoice:zz"),
                ("error", "خطأ ما")]
        stamps = [f"2026-10-03 09:0{i}" for i in range(len(msgs))]
        return ce.build_entries(msgs, stamps, image_of={"a": jpeg((10, 120, 80))}.get, doc_label=lambda r: "فاتورة بيع INV-1")

    def test_image_and_default_caption_merge_into_one_user_entry(self):
        es = self.entries()
        self.assertEqual([e.role for e in es], ["user", "assistant", "user", "assistant", "assistant", "assistant", "system"])
        self.assertEqual(len(es[0].images), 1)
        self.assertEqual(es[0].text, "")                       # «📷 صورة مرفقة» الافتراضي حُذف
        self.assertEqual(es[0].time, "2026-10-03 09:00")
        self.assertEqual(es[3].kind, "pictures")
        self.assertEqual(len(es[3].images), 1)
        self.assertIn("INV-1", es[5].text)

    def test_plain_text_strips_markdown_and_points_to_table_sheet(self):
        txt = ce.plain_text(self.MD, ["جدول 1"])
        self.assertNotIn("**", txt)
        self.assertNotIn("##", txt)
        self.assertIn("[جدول «جدول 1»", txt)
        self.assertNotIn("|--", txt)
        self.assertIn("• موجود بالمخزون", txt)

    def test_workbook_has_conversation_sheet_and_one_sheet_per_table(self):
        data = ce.chat_workbook(self.entries(), "العمر")
        wb = load_workbook(io.BytesIO(data))
        self.assertEqual(wb.sheetnames, ["المحادثة", "جدول 1"])
        chat = sheet_values(data, "المحادثة")
        self.assertEqual(chat[0], ["#", "الوقت", "المرسل", "الرسالة"])
        self.assertEqual(len(chat), 1 + 7)
        self.assertEqual(chat[1][2], "أنت")
        self.assertIn("1 صورة مرفقة", chat[1][3])
        self.assertEqual(sheet_values(data, "جدول 1"), [["الصنف", "الكمية", "السعر"], ["فلتر", 6, 20], ["زيت", 30, 3]])
        self.assertTrue(wb["المحادثة"].sheet_view.rightToLeft)
        self.assertEqual(wb["المحادثة"].column_dimensions["D"].width, 110)

    def test_workbook_without_tables_has_only_conversation_sheet(self):
        data = ce.chat_workbook(ce.build_entries([("user", "سؤال"), ("assistant", "جواب")], ["a", "b"]))
        self.assertEqual(load_workbook(io.BytesIO(data)).sheetnames, ["المحادثة"])

    def test_pdf_doc_renders_content_and_page_numbers(self):
        doc = ce.chat_doc(self.entries(), {"businessName": "العمر"})
        self.assertGreaterEqual(len(doc.pages), 1)
        self.assertTrue(all(p.getbbox() is not None and len(set(p.convert("L").resize((40, 40)).getdata())) > 3 for p in doc.pages))
        pdf = io.BytesIO()
        doc.pages[0].save(pdf, "PDF", save_all=True, append_images=doc.pages[1:])
        self.assertTrue(pdf.getvalue().startswith(b"%PDF"))

    def test_long_conversation_paginates_without_crashing(self):
        msgs = []
        for i in range(40):
            msgs += [("user", f"سؤال رقم {i} " + "كلمة " * 30), ("assistant", self.MD + "\n" + "نص طويل جداً " * 80)]
        doc = ce.chat_doc(ce.build_entries(msgs, [""] * len(msgs)), {})
        self.assertGreater(len(doc.pages), 5)

    def test_empty_conversation_still_makes_a_document(self):
        self.assertEqual(len(ce.chat_doc([], {}).pages), 1)

    def test_emoji_removed_for_pdf_font(self):
        self.assertEqual(ce._pdf_text("تم ✅ 📷 بنجاح"), "تم بنجاح")

    def test_header_is_not_orphaned_at_page_bottom(self):
        """رد المساعد وصوره يبدآن معاً: إن لم تتسع الصور تنتقل الترويسة معها."""
        from daftari.printing import render as R
        doc = R.Doc()
        doc.y = doc.h - doc.margin - 120               # لا يتسع إلا لترويسة صغيرة
        n = len(doc.pages)
        e = ce.ChatEntry("assistant", ce.BOT_LABEL, "صور", "t", kind="pictures", images=[("a", jpeg())])
        ce._draw_entry(doc, R, e)
        self.assertEqual(len(doc.pages), n + 1)
        self.assertLess(doc.y, doc.h // 2)


class ChatPrintHeader(S.Smoke):
    def test_logo_asset_exists_and_header_follows_settings(self):
        from daftari.printing import render as R
        self.assertTrue((R.ASSETS / ce.CHAT_LOGO).exists())
        es = ce.build_entries([("user", "س"), ("assistant", "ج")], ["", ""])
        default = ce.chat_doc(es, {}).pages[0].tobytes()
        custom = ce.chat_doc(es, {"chatPrintNote": "عبارة مخصصة للتدقيق", "chatPrintAiType": "نموذج آخر"}).pages[0].tobytes()
        hidden = ce.chat_doc(es, {"chatPrintNote": "", "chatPrintAiType": ""}).pages[0].tobytes()
        self.assertEqual(len({default, custom, hidden}), 3)

    def _logo_on_gradient(self):
        """شعار داكن على خلفية رمادية فاتحة متدرجة، بفجوة فاتحة داخله (كمعدن لامع)."""
        from PIL import ImageDraw
        im = Image.new("RGB", (300, 240))
        px = im.load()
        for x in range(300):
            for y in range(240):
                v = 250 - int(18 * x / 300)
                px[x, y] = (v, v, v)
        d = ImageDraw.Draw(im)
        d.rectangle((90, 60, 210, 180), fill=(30, 30, 30))
        d.rectangle((130, 100, 170, 140), fill=(225, 225, 225))      # لمعة داخلية فاتحة
        b = io.BytesIO(); im.save(b, "JPEG", quality=95)
        return b.getvalue()

    def test_clean_logo_removes_background_keeps_inner_highlight_and_crops(self):
        png = ce.clean_logo(self._logo_on_gradient())
        im = Image.open(io.BytesIO(png))
        self.assertEqual(im.mode, "RGBA")
        self.assertLess(im.width, 160)                                   # قُصّ على حدود الشعار
        a = im.getchannel("A")
        self.assertEqual(a.getpixel((2, 2)), 0)                          # الزاوية شفافة
        self.assertEqual(a.getpixel((im.width // 2, im.height // 2)), 255)   # اللمعة الداخلية مصمتة لا ثقب

    def test_transparent_png_is_kept_and_blank_image_rejected(self):
        t = Image.new("RGBA", (100, 100), (0, 0, 0, 0))
        t.paste((200, 0, 0, 255), (30, 30, 70, 70))
        b = io.BytesIO(); t.save(b, "PNG")
        out = Image.open(io.BytesIO(ce.clean_logo(b.getvalue())))
        self.assertLess(out.width, 60)
        blank = io.BytesIO(); Image.new("RGB", (50, 50), (255, 255, 255)).save(blank, "PNG")
        with self.assertRaises(ValueError):
            ce.clean_logo(blank.getvalue())
        with self.assertRaises(ValueError):
            ce.clean_logo(b"not an image")

    def test_uploaded_logo_overrides_default_and_reset_restores(self):
        import base64
        default = ce.logo_bytes({})
        self.assertTrue(default)
        mine = ce.clean_logo(self._logo_on_gradient())
        self.assertEqual(ce.logo_bytes({"chatPrintLogoB64": base64.b64encode(mine).decode()}), mine)
        self.assertEqual(ce.logo_bytes({"chatPrintLogoB64": "", "chatPrintLogo": ""}), default)
        es = ce.build_entries([("user", "س")], [""])
        pages = {k: ce.chat_doc(es, st).pages[0].tobytes() for k, st in
                 {"d": {}, "m": {"chatPrintLogoB64": base64.b64encode(mine).decode()}}.items()}
        self.assertNotEqual(pages["d"], pages["m"])

    async def test_settings_upload_saves_transparent_logo_and_reset(self):
        self.app.tab, self.app.sub = "settings", {}
        ctl = await self.app._build_current()
        async def fake_pick(extensions=None, images=False):
            return ("logo.jpg", self._logo_on_gradient())
        self.app.pick_file_bytes = fake_pick
        await S.find_btn(ctl, "رفع شعار").on_click(None)
        st = self.led.settings
        self.assertTrue(st["chatPrintLogoB64"])
        self.assertTrue(Path(st["chatPrintLogo"]).exists())
        import base64
        self.assertEqual(Image.open(io.BytesIO(base64.b64decode(st["chatPrintLogoB64"]))).mode, "RGBA")
        await S.find_btn(ctl, "الافتراضي").on_click(None)
        self.assertEqual(self.led.settings["chatPrintLogoB64"], "")

    async def test_settings_upload_rejects_bad_file_with_message(self):
        self.app.tab, self.app.sub = "settings", {}
        ctl = await self.app._build_current()
        async def bad(extensions=None, images=False):
            return ("x.png", b"garbage")
        self.app.pick_file_bytes = bad
        await S.find_btn(ctl, "رفع شعار").on_click(None)
        self.assertFalse(self.led.settings.get("chatPrintLogoB64"))
        self.assertIn("تعذّرت قراءة الصورة", last_toast(self.page))

    async def test_settings_card_edits_print_header_and_persists(self):
        self.app.tab, self.app.sub = "settings", {}
        ctl = await self.app._build_current()
        self.assertIn("ترويسة طباعة المحادثة", [c.args[0] for c in S.walk(ctl) if getattr(c, "args", None) and isinstance(c.args[0], str)])
        self.led.settings["chatPrintNote"] = "نص جديد"
        await self.led.repo.save_settings(self.led.settings)
        self.assertEqual(ce.chat_doc([], self.led.settings).pages[0].size, (1240, 1754))


# ===================================================================================== الحفظ
class SaveBytesTests(S.Smoke):
    async def asyncSetUp(self):
        await super().asyncSetUp()
        self.out = Path(self.tmp.name) / "chosen"
        self.out.mkdir()

    async def test_desktop_picker_returns_path_but_file_is_written_by_us(self):
        self.app.picker.save_path = str(self.out / "تقرير.xlsx")
        where = await self.app.save_bytes("تقرير.xlsx", b"DATA")
        self.assertEqual(where, str(self.out / "تقرير.xlsx"))
        self.assertEqual((self.out / "تقرير.xlsx").read_bytes(), b"DATA")

    async def test_cancel_is_not_reported_as_saved(self):
        self.app.picker.save_path = False
        self.assertIsNone(await self.app.save_bytes("a.xlsx", b"DATA"))
        self.assertFalse(await self.app.save_with_toast("a.xlsx", b"DATA", "الملف"))
        self.assertIn("أُلغي", last_toast(self.page))
        self.assertEqual(list(self.out.iterdir()), [])

    async def test_missing_extension_is_added(self):
        self.app.picker.save_path = str(self.out / "جدولي")
        where = await self.app.save_bytes("جدولي.xlsx", b"DATA")
        self.assertTrue(where.endswith("جدولي.xlsx"))
        self.assertEqual([p.name for p in self.out.iterdir()], ["جدولي.xlsx"])

    async def test_extension_less_file_already_written_by_picker_is_renamed_not_duplicated(self):
        (self.out / "جدولي").write_bytes(b"DATA")                # بعض المنصات تكتب الملف بنفسها
        self.app.picker.save_path = str(self.out / "جدولي")
        await self.app.save_bytes("جدولي.xlsx", b"DATA")
        self.assertEqual([p.name for p in self.out.iterdir()], ["جدولي.xlsx"])

    async def test_dialog_receives_extension_filter(self):
        self.app.picker.save_path = str(self.out / "f.xlsx")
        await self.app.save_bytes("f.xlsx", b"x")
        kw = self.app.picker.save_calls[-1]
        self.assertEqual(kw["allowed_extensions"], ["xlsx"])
        self.assertEqual(kw["src_bytes"], b"x")

    async def test_unwritable_target_falls_back_to_app_folder_with_explanation(self):
        self.app.picker.save_path = str(self.out / "no-such-dir" / "f.xlsx")
        where = await self.app.save_bytes("f.xlsx", b"DATA")
        self.assertIn("تعذّر الحفظ في المكان المختار", where)
        self.assertEqual((self.app.work_dir / "f.xlsx").read_bytes(), b"DATA")

    async def test_unsupported_platform_falls_back_to_app_folder(self):
        self.app.picker.raise_on_save = RuntimeError("unsupported")
        where = await self.app.save_bytes("f.xlsx", b"DATA")
        self.assertEqual(where, str(self.app.work_dir / "f.xlsx"))

    async def test_web_mode_lets_browser_download_and_never_writes_server_disk(self):
        self.page.web = True
        self.app.picker.save_path = False                       # في الويب None شائعة ولا تعني الإلغاء
        where = await self.app.save_bytes("f.xlsx", b"DATA")
        self.assertTrue(where)
        self.assertFalse((self.app.work_dir / "f.xlsx").exists())
        self.app.picker.raise_on_save = RuntimeError("boom")
        self.assertFalse(await self.app.save_with_toast("f.xlsx", b"DATA"))
        self.assertIn("تعذّر الحفظ", last_toast(self.page))
        self.assertFalse((self.app.work_dir / "f.xlsx").exists())

    async def test_open_file_is_disabled_on_web(self):
        self.page.web = True
        self.assertFalse(self.app.open_file(self.out))

    async def test_success_toast_names_the_location(self):
        self.app.picker.save_path = str(self.out / "f.xlsx")
        self.assertTrue(await self.app.save_with_toast("f.xlsx", b"x", "ملف Excel"))
        self.assertIn(str(self.out / "f.xlsx"), last_toast(self.page))


# ===================================================================================== صور المساعد + أزرار الطباعة
class AssistantImagesAndPrinting(S.Smoke):
    async def asyncSetUp(self):
        await super().asyncSetUp()
        self.item = next(i for i in self.led.inventory if i.name == "فلتر زيت تويوتا")
        self.other = next(i for i in self.led.inventory if i.name == "زيت محرك 5W30")
        await self.app.images.put(self.item.id, jpeg((20, 140, 90)))

    def prime(self, *responses):
        cfg = GeminiConfig(api_key="K")
        tools = ToolBox(self.led, None, self.app.home, assistant_tab._confirm_factory(self.app),
                        lambda: self.app.ai_state["auto"], images=self.app.images)
        self.app.ai_state.update(messages=[], stamps=[], busy=False, auto=False, cfg=cfg,
                                 assistant=Assistant(GeminiClient(cfg, FakeTransport(*responses)), tools, "المتجر"))

    async def send(self, ctl, text):
        tf = next(c for c in S.walk(ctl) if getattr(c, "kw", {}).get("label", "").startswith("اكتب سؤالك"))
        tf.value = text
        await S.find_btn(ctl, "إرسال").on_click(None)

    def texts(self, ctl):
        return [c.args[0] for c in S.walk(ctl) if getattr(c, "args", None) and isinstance(c.args[0], str)]

    # ---- الأداة
    async def test_tool_returns_only_items_with_pictures_and_records_gallery(self):
        tb = ToolBox(self.led, None, self.app.home, images=self.app.images)
        out = await tb.execute("show_item_images", {"query": ""})
        self.assertFalse(out["ok"])                                         # لا بد من تحديد الصنف
        out = await tb.execute("show_item_images", {"query": "فلتر زيت تويوتا"})
        self.assertTrue(out["ok"], out)
        self.assertEqual([x["name"] for x in out["result"]["shown"]], ["فلتر زيت تويوتا"])
        self.assertNotIn("id", out["result"]["shown"][0])                    # النموذج لا يحتاج المعرّف
        self.assertEqual(len(tb.pictures), 1)
        self.assertEqual(tb.pictures[0][0]["id"], self.item.id)

    async def test_tool_skips_items_without_pictures_and_reports_them(self):
        tb = ToolBox(self.led, None, self.app.home, images=self.app.images)
        out = await tb.execute("show_item_images", {"item_ids": f"{self.item.id}, {self.other.id}"})
        self.assertTrue(out["ok"], out)
        self.assertEqual(len(out["result"]["shown"]), 1)
        self.assertEqual(out["result"]["matches_without_image"], ["زيت محرك 5W30"])

    async def test_tool_errors_are_readable(self):
        tb = ToolBox(self.led, None, self.app.home, images=self.app.images)
        no_pic = await tb.execute("show_item_images", {"query": "زيت محرك"})
        self.assertFalse(no_pic["ok"])
        self.assertIn("لا توجد صورة", no_pic["error"])
        none = await tb.execute("show_item_images", {"query": "صنف غير موجود أبداً"})
        self.assertFalse(none["ok"])
        self.assertEqual(tb.pictures, [])

    async def test_tool_creates_its_own_image_store_when_not_injected(self):
        tb = ToolBox(self.led, None, self.app.home)                          # app.home/images هو نفس مجلد التطبيق
        out = await tb.execute("show_item_images", {"query": "فلتر زيت تويوتا"})
        self.assertTrue(out["ok"], out)

    async def test_limit_goes_from_one_to_one_hundred(self):
        for i in range(30):
            it = await self.led.add_item(f"مكنسة {i}", f"BR{i}", 1, 2, 1)
            await self.app.images.put(it.id, jpeg())
        tb = ToolBox(self.led, None, self.app.home, images=self.app.images)
        for limit, expect in ((1, 1), (5, 5), (50, 30), (500, 30)):          # لا سقف 8: حتى كل المطابقات (الأقصى 100)
            out = await tb.execute("show_item_images", {"query": "مكنسة", "limit": limit})
            self.assertEqual(len(out["result"]["shown"]), expect, limit)
        out = await tb.execute("show_item_images", {"query": "مكنسة"})
        self.assertEqual(len(out["result"]["shown"]), 12)                      # الافتراضي 12 إن لم يحدّد النموذج

    # ---- الواجهة
    async def test_assistant_sends_pictures_into_the_chat(self):
        self.prime(reply_json([{"functionCall": {"name": "show_item_images", "args": {"query": "فلتر زيت تويوتا"}}}]),
                   reply_json([{"text": "هذه صورة الفلتر."}]))
        self.app.tab, self.app.sub = "assistant", {}
        ctl = await self.app._build_current()
        await self.send(ctl, "أرني صورة الفلتر")
        st = self.app.ai_state
        self.assertEqual([m[0] for m in st["messages"]], ["user", "assistant", "pictures"])
        self.assertEqual(len(st["stamps"]), 3)
        self.assertTrue(all(st["stamps"]))
        t = self.texts(await self.app._build_current())
        self.assertIn("فلتر زيت تويوتا", t)
        self.assertIn("صور من المخزون — اضغط على الصورة لتكبيرها", t)

    async def test_picture_tile_opens_zoom_dialog(self):
        self.prime(reply_json([{"functionCall": {"name": "show_item_images", "args": {"query": "فلتر زيت تويوتا"}}}]),
                   reply_json([{"text": "تفضل."}]))
        self.app.tab, self.app.sub = "assistant", {}
        ctl = await self.app._build_current()
        await self.send(ctl, "صورة")
        ctl = await self.app._build_current()
        tile = next(c for c in S.walk(ctl) if getattr(c, "tooltip", None) == "اضغط لعرض صورة المنتج")
        tile.on_click(None)
        self.assertTrue(self.page.dialogs)

    async def test_system_prompt_tells_model_about_images(self):
        self.prime(reply_json([{"text": "x"}]))
        self.assertIn("show_item_images", self.app.ai_state["assistant"].system)

    # ---- طباعة المحادثة
    async def chat(self):
        self.prime(reply_json([{"text": "## ملخص\n| الصنف | الكمية |\n|--|--|\n| فلتر | 6 |"}]))
        self.app.tab, self.app.sub = "assistant", {}
        ctl = await self.app._build_current()
        await self.send(ctl, "اعرض المخزون")
        return await self.app._build_current()

    async def test_print_conversation_opens_pdf_document_and_back_returns(self):
        ctl = await self.chat()
        btn = next(c for c in S.walk(ctl) if getattr(c, "kw", {}).get("tooltip") == "طباعة المحادثة (PDF)")
        await btn.on_click(None)
        self.assertIn("doc", self.app.sub)
        self.assertGreaterEqual(len(self.app.sub["pages"]), 1)
        self.assertTrue(self.app.sub["pages"][0].startswith(b"\x89PNG"))
        self.assertIn("محادثة-المساعد", self.app.sub["filename"])
        self.assertEqual(self.app.doc_return, ("assistant", {}))
        self.assertEqual(len(self.app.ai_state["messages"]), 2)              # المحادثة محفوظة عند الرجوع

    async def test_export_conversation_to_excel_saves_real_workbook(self):
        ctl = await self.chat()
        out = Path(self.tmp.name) / "o.xlsx"
        self.app.picker.save_path = str(out)
        btn = next(c for c in S.walk(ctl) if getattr(c, "kw", {}).get("tooltip") == "تصدير المحادثة إلى Excel")
        await btn.on_click(None)
        wb = load_workbook(out)
        self.assertEqual(wb.sheetnames, ["المحادثة", "ملخص"])       # اسم الورقة من عنوان الجدول
        self.assertEqual(self.app.picker.save_calls[-1]["file_name"].endswith(".xlsx"), True)
        self.assertIn(str(out), last_toast(self.page))

    async def test_export_buttons_warn_when_chat_is_empty(self):
        self.prime(reply_json([{"text": "x"}]))
        self.app.tab, self.app.sub = "assistant", {}
        ctl = await self.app._build_current()
        for tip in ("طباعة المحادثة (PDF)", "تصدير المحادثة إلى Excel"):
            await next(c for c in S.walk(ctl) if getattr(c, "kw", {}).get("tooltip") == tip).on_click(None)
            self.assertIn("لا توجد رسائل", last_toast(self.page))
        self.assertNotIn("doc", self.app.sub)

    async def test_single_reply_has_its_own_print_button(self):
        ctl = await self.chat()
        t = self.texts(ctl)
        self.assertIn("طباعة / PDF", t)
        self.assertIn("تصدير الجدول إلى Excel", t)
        await S.find_btn(ctl, "طباعة / PDF").on_click(None)
        self.assertIn("doc", self.app.sub)
        self.assertIn("رد-المساعد", self.app.sub["filename"])

    async def test_clear_resets_timestamps_too(self):
        ctl = await self.chat()
        self.assertEqual(len(self.app.ai_state["stamps"]), 2)
        next(c for c in S.walk(ctl) if getattr(c, "kw", {}).get("tooltip") == "محادثة جديدة").on_click(None)
        self.assertEqual(self.app.ai_state["messages"], [])
        self.assertEqual(self.app.ai_state["stamps"], [])

    async def test_chat_state_without_stamps_key_still_works(self):
        """الاختبارات القديمة والجلسات الجارية قد تفتقد مفتاح stamps."""
        self.prime(reply_json([{"text": "ok"}]))
        self.app.ai_state.pop("stamps", None)
        self.app.tab, self.app.sub = "assistant", {}
        ctl = await self.app._build_current()
        await self.send(ctl, "سؤال")
        self.assertEqual(len(self.app.ai_state["stamps"]), 2)

    # ---- بطاقة الملف
    async def test_file_card_save_button_writes_the_exported_file(self):
        self.prime(reply_json([{"functionCall": {"name": "export_dataset", "args": {"dataset": "inventory"}}}]),
                   reply_json([{"text": "جاهز."}]))
        self.app.tab, self.app.sub = "assistant", {}
        ctl = await self.app._build_current()
        await self.send(ctl, "صدّر")
        ctl = await self.app._build_current()
        out = Path(self.tmp.name) / "saved.xlsx"
        self.app.picker.save_path = str(out)
        await S.find_btn(ctl, "حفظ باسم…").on_click(None)
        self.assertEqual(load_workbook(out).sheetnames, ["المخزون"])

    async def test_file_card_on_web_downloads_instead_of_opening_on_server(self):
        self.page.web = True
        self.prime(reply_json([{"functionCall": {"name": "export_dataset", "args": {"dataset": "inventory"}}}]),
                   reply_json([{"text": "جاهز."}]))
        self.app.tab, self.app.sub = "assistant", {}
        ctl = await self.app._build_current()
        await self.send(ctl, "صدّر")
        ctl = await self.app._build_current()
        self.assertIn("تنزيل", self.texts(ctl))
        self.assertNotIn("فتح", self.texts(ctl))
        await S.find_btn(ctl, "تنزيل").on_click(None)
        self.assertTrue(self.app.picker.save_calls)
        self.assertTrue(self.app.picker.save_calls[-1]["src_bytes"].startswith(b"PK"))   # xlsx = zip


if __name__ == "__main__":
    unittest.main()


class MessagePrintStyle(unittest.TestCase):
    """طباعة رد واحد: الرسالة نفسها كما على الشاشة (شبكة جدول كاملة + الشعار والكلام بجانبه) بلا شريط اسم المرسل."""
    MD = "### العملاء\n\n| العميل | المنطقة | الرصيد |\n|---|---|---|\n| **أحمد** | — | $10.00 |\n| سعيد | الباب | $5.00 |\n"

    def _doc(self, style):
        entries = ce.build_entries([("assistant", self.MD)], ["٣ أكتوبر ٢٠٢٦ 04:12"])
        return ce.chat_doc(entries, {"businessName": "شركة العمر"}, "رد مساعد شركة العمر", style)

    def test_message_style_is_more_compact_and_has_no_black_band(self):
        msg, chat = self._doc("message"), self._doc("chat")
        self.assertEqual(len(msg.pages), 1)
        self.assertLess(msg.y, chat.y)                       # بلا شريط المرسل ولا عنوان كبير
        # أول صف بعد الترويسة ليس شريطاً أسود كما في وضع المحادثة
        im = msg.pages[0].convert("RGB")
        black_rows = sum(1 for y in range(380, 700, 4)
                         if all(sum(im.getpixel((x, y))) < 120 for x in range(150, 1100, 60)))
        self.assertEqual(black_rows, 0)

    def test_tables_have_full_grid_in_both_styles(self):
        for style in ("message", "chat"):
            im = self._doc(style).pages[0].convert("RGB")
            grey_cols = 0                                    # خطوط عمودية رمادية (حدود الخلايا) في منطقة الجدول
            for x in range(70, 1170):
                run = sum(1 for y in range(380, 900) if im.getpixel((x, y)) == (150, 150, 150))
                if run > 60:
                    grey_cols += 1
            self.assertGreaterEqual(grey_cols, 4, style)     # 3 أعمدة = 4 حدود عمودية على الأقل (قد تكون سميكة)

    def test_default_style_unchanged_signature(self):
        entries = ce.build_entries([("user", "سؤال"), ("assistant", "جواب")], ["", ""])
        self.assertTrue(ce.chat_doc(entries, {}, "x").pages)


class ReplyImageExport(unittest.TestCase):
    """إصدار صور الرد (فلتر…) مع الطباعة والإكسل + زر الواجهة + تعليمات الجداول للمساعد."""
    MD = ("### بطاقة الصنف\n\n| الحقل | القيمة |\n|---|---|\n| الماركة | MANN |\n| القطر الخارجي | 76 مم |\n\n"
          "### المرجعيات\n\n| الماركة | الرقم |\n|---|---|\n| Bosch | 0 986 AF1 |\n")

    @staticmethod
    def _jpeg(color=(30, 90, 160), size=(900, 700)):
        b = io.BytesIO(); Image.new("RGB", size, color).save(b, "JPEG"); return b.getvalue()

    def test_excel_embeds_images_beside_first_table(self):
        from daftari.ai import excel
        from openpyxl import load_workbook
        imgs = [("HU 7008 z", self._jpeg()), ("OC 1234", self._jpeg((160, 40, 40)))]
        data = excel.build_workbook(excel.parse_markdown_tables(self.MD), "شركة العمر", imgs)
        wb = load_workbook(io.BytesIO(data))
        first = wb.worksheets[0]
        self.assertEqual(len(first._images), 2)                       # الصورتان داخل الملف فعلاً
        col = first._images[0].anchor._from.col + 1
        self.assertGreater(col, 2)                                    # بعد أعمدة الجدول (عمود فارغ بينهما)
        self.assertEqual(first.cell(row=1, column=col).value, "HU 7008 z")
        self.assertEqual(len(wb.worksheets[1]._images), 0)            # الجدول الثاني بلا صور
        self.assertEqual(first.cell(row=1, column=1).value, "الحقل")  # الجدول سليم

    def test_excel_images_only_and_bad_image(self):
        from daftari.ai import excel
        from openpyxl import load_workbook
        data = excel.build_workbook([], "", [("x", self._jpeg()), ("تالفة", b"not-an-image")])
        ws = load_workbook(io.BytesIO(data)).worksheets[0]
        self.assertEqual((ws.title, len(ws._images)), ("الصور", 1))   # التالفة تُتجاوز
        with self.assertRaises(ValueError):
            excel.build_workbook([], "")
        # التوافق: بلا صور يعمل كما كان
        self.assertTrue(excel.build_workbook(excel.parse_markdown_tables(self.MD)))

    def test_pdf_message_style_draws_large_image_above_table(self):
        entries = ce.build_entries([("assistant", self.MD)], ["٣ أكتوبر ٢٠٢٦ 04:12"])
        plain = ce.chat_doc(entries, {}, "رد", "message")
        entries[0].images = [("HU 7008 z", self._jpeg((200, 30, 30)))]
        with_img = ce.chat_doc(entries, {}, "رد", "message")
        self.assertGreater(with_img.y, plain.y + 250)                 # صورة بحجم 300 بكسل أخذت مكاناً
        pg = with_img.pages[0].convert("RGB")
        reds = sum(1 for x in range(100, 1150, 25) for y in range(300, 900, 25)
                   if (lambda c: c[0] > 150 and c[1] < 90 and c[2] < 90)(pg.getpixel((x, y))))     # JPEG يغيّر اللون قليلاً
        self.assertGreater(reds, 20)                                  # الصورة ظهرت فعلاً في الصفحة

    def test_ui_button_only_when_reply_has_images(self):
        from daftari.ui.tabs import assistant as A
        from daftari.tests import fake_flet

        class App:
            page = fake_flet.Page()
        calls = []
        orig = A._image_export_button
        A._image_export_button = lambda *a, **k: (calls.append(a[4]), orig(*a, **k))[1]
        try:
            A._bubble(App(), "assistant", self.MD, "t")
            self.assertEqual(calls, [])
            A._bubble(App(), "assistant", self.MD, "t", [("x", self._jpeg())])
            self.assertEqual(len(calls), 1)
        finally:
            A._image_export_button = orig

    def test_prompt_demands_tables_and_full_specs(self):
        from daftari.ai.agent import Assistant
        sysmsg = Assistant(None, None, "متجر").system
        for needle in ("جداول Markdown", "القطر الخارجي", "المادة المصنوع منها", "المرجعيات المقابلة", "غير متوفر", "google_search"):
            self.assertIn(needle, sysmsg)
        self.assertEqual(Assistant.MAX_STEPS, 20)


class ManyImages(unittest.TestCase):
    """لا حد 8: من صورة واحدة حتى 100 صورة تخرج مع الطباعة والإكسل ومع بحث الكتالوج."""
    @staticmethod
    def _imgs(n):
        out = []
        for i in range(n):
            b = io.BytesIO(); Image.new("RGB", (800, 600), (i * 2 % 256, 80, 160)).save(b, "JPEG")
            out.append((f"صنف {i + 1}", b.getvalue()))
        return out

    def test_excel_100_images_grid(self):
        from daftari.ai import excel
        from openpyxl import load_workbook
        md = "| الحقل | القيمة |\n|---|---|\n| أ | 1 |\n"
        for n in (1, 6, 7, 100):
            data = excel.build_workbook(excel.parse_markdown_tables(md), "", self._imgs(n))
            ws = load_workbook(io.BytesIO(data)).worksheets[0]
            self.assertEqual(len(ws._images), n, n)
            cells = {(im.anchor._from.row, im.anchor._from.col) for im in ws._images}
            self.assertEqual(len(cells), n, "كل صورة بمكان مختلف (لا تتراكب)")
        self.assertLess(len(data), 3_000_000)                          # 100 صورة لا تضخّم الملف

    def test_pdf_100_images_all_drawn_across_pages(self):
        entries = ce.build_entries([("assistant", "### جدول\n\n| أ | ب |\n|---|---|\n| 1 | 2 |\n")], [""])
        for n in (1, 5, 40, 100):
            entries[0].images = self._imgs(n)
            doc = ce.chat_doc(entries, {}, "رد", "message")
            drawn = 0
            for pg in doc.pages:
                im = pg.convert("RGB")
                for y in range(0, im.height, 30):
                    for x in range(0, im.width, 30):
                        c = im.getpixel((x, y))
                        if c[2] > 140 and c[1] < 120 and c[0] < 200:
                            drawn += 1
            self.assertGreater(drawn, 0, n)
            if n == 100:
                self.assertGreater(len(doc.pages), 2)                  # امتدت على عدة صفحات بدل أن تُقتطع
        self.assertEqual(ce._img_size(2, False), 300)
        self.assertLess(ce._img_size(100, False), ce._img_size(10, False))

    def test_reply_cap_is_100_and_thumbs_are_light(self):
        from daftari.ui.tabs import assistant as A
        self.assertEqual(A.MAX_REPLY_IMAGES, 100)
        big = io.BytesIO(); Image.new("RGB", (1000, 1000), (10, 120, 200)).save(big, "JPEG", quality=95)
        th = A._thumb_bytes(big.getvalue())
        self.assertLess(len(th), len(big.getvalue()))
        self.assertLessEqual(max(Image.open(io.BytesIO(th)).size), A.THUMB_SIDE)
