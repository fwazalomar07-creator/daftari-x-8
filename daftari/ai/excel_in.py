"""قراءة ملف Excel/CSV يرسله المالك للمساعد → جدول بسيط (عناوين + صفوف) ليسجّل منه أصنافاً."""
from __future__ import annotations

import csv
import io

MAX_ROWS = 3000


def _clean(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def parse_table(name: str, raw: bytes) -> dict:
    """يرجع {"sheet", "headers", "rows"}؛ يرفع ValueError برسالة عربية إن تعذّرت القراءة."""
    ext = name.lower().rsplit(".", 1)[-1] if "." in name else ""
    rows: list[list[str]] = []
    sheet = ""
    if ext in ("xlsx", "xlsm"):
        from openpyxl import load_workbook
        try:
            wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        except Exception as e:  # noqa: BLE001
            raise ValueError("تعذّرت قراءة ملف Excel — تأكد أنه .xlsx سليم") from e
        ws = wb.worksheets[0]
        sheet = ws.title
        for r in ws.iter_rows(values_only=True):
            line = [_clean(c) for c in r]
            if any(line):
                rows.append(line)
            if len(rows) > MAX_ROWS:
                break
    elif ext in ("csv", "txt"):
        text = raw.decode("utf-8-sig", errors="replace")
        for r in csv.reader(io.StringIO(text)):
            line = [_clean(c) for c in r]
            if any(line):
                rows.append(line)
            if len(rows) > MAX_ROWS:
                break
    elif ext == "xls":
        raise ValueError("صيغة .xls القديمة غير مدعومة — احفظ الملف من Excel بصيغة .xlsx ثم أعد رفعه")
    else:
        raise ValueError("اختر ملف Excel (.xlsx) أو CSV أو صورة")
    if not rows:
        raise ValueError("الملف فارغ")
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    headers = [h or f"عمود {i + 1}" for i, h in enumerate(rows[0])]
    return {"sheet": sheet, "headers": headers, "rows": rows[1:]}
