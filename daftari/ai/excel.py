"""تصدير الجداول إلى Excel (.xlsx): من جداول Markdown التي يكتبها المساعد، أو من جداول يبنيها البرنامج مباشرة.

الشكل: ورقة من اليمين لليسار (RTL)، ترويسة خضراء بخط أبيض عريض، صفوف متبادلة اللون، تجميد الترويسة، فلتر تلقائي،
عرض أعمدة تلقائي، والأرقام تُكتب كأرقام حقيقية (لا نصوص) بتنسيق دولار/نسبة/عدد حتى تعمل عليها المعادلات والفرز.
"""
from __future__ import annotations

import io
import json
import math
import re
from dataclasses import dataclass, field
from datetime import datetime
from decimal import Decimal
from typing import Any, Sequence

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..core.money import normalize_digits

MONEY_FMT = '"$"#,##0.00'
INT_FMT = "#,##0"
DEC_FMT = "#,##0.00"
PCT_FMT = "0.0%"
FONT = "Arial"                      # يدعم العربية ومتوفر في ويندوز وماك وأوفيس
GREEN, MINT, LINE = "059669", "ECFDF5", "E5E7EB"
_BAD_SHEET = re.compile(r"[\[\]:*?/\\]")
_BAD_FILE = re.compile(r'[\\/:*?"<>|\r\n]+')
MAX_CELL_CHARS = 32767               # حد إكسل لخلية واحدة؛ ما زاد عنه يجعل الملف «تالفاً» عند الفتح
RESERVED_SHEETS = {"history"}        # اسم محجوز في إكسل


@dataclass
class Sheet:
    name: str
    headers: list[str]
    rows: list[list[Any]]
    money_cols: set[int] = field(default_factory=set)     # أعمدة تُنسَّق بالدولار
    totals: bool = False                                  # صف «الإجمالي» بمعادلات SUM لأعمدة الدولار
    widths: dict[int, float] = field(default_factory=dict)  # عرض ثابت لأعمدة محددة (يغلب العرض التلقائي)


# ---------------------------------------------------------------------------------------------- تحويل الخلايا
_NUM = re.compile(r"^(-?)\s*(\$?)\s*(-?)(\d[\d,]*(?:\.\d+)?|\.\d+)\s*(\$?)$")
_PCT = re.compile(r"^(-?\d+(?:\.\d+)?)\s*%$")
_MD = re.compile(r"(\*\*|__|`|~~)")


def clean_text(s: Any) -> str:
    t = _MD.sub("", str(s)).replace("<br>", "\n").replace("<br/>", "\n")
    return ILLEGAL_CHARACTERS_RE.sub("", t).strip()      # محارف التحكم (\x00-\x1f) تُسقط openpyxl بخطأ


def _fit(text: str) -> str:
    return text if len(text) <= MAX_CELL_CHARS else text[: MAX_CELL_CHARS - 1] + "…"


def coerce(value: Any, col: int = 0, money_cols: set[int] | None = None) -> tuple[Any, str | None]:
    """يرجع (القيمة، تنسيق الرقم). نصوص تشبه الأرقام تتحول لأرقام، إلا الأكواد والهواتف (أصفار بادئة / أرقام طويلة)."""
    money_cols = money_cols or set()
    if value is None:
        return None, None
    if isinstance(value, bool):
        return ("نعم" if value else "لا"), None
    if isinstance(value, (int, Decimal, float)):
        n = float(value) if not isinstance(value, int) else value
        if isinstance(n, float) and not math.isfinite(n):      # NaN/Infinity تُفسد الملف في إكسل
            return None, None
        if col in money_cols:
            return n, MONEY_FMT
        return n, (INT_FMT if isinstance(value, int) or float(value).is_integer() else DEC_FMT)
    text = _fit(clean_text(value))
    if not text:
        return None, None
    norm = normalize_digits(text)
    m = _PCT.match(norm)
    if m:
        return float(m.group(1)) / 100, PCT_FMT
    m = _NUM.match(norm)
    if m:
        sign1, dollar1, sign2, digits, dollar2 = m.groups()
        raw = digits.replace(",", "")
        is_code = (raw.startswith("0") and len(raw) > 1 and "." not in raw) or len(raw.split(".")[0]) >= 9
        if not is_code:
            n = float(raw)
            if sign1 or sign2:
                n = -n
            if dollar1 or dollar2 or col in money_cols:
                return n, MONEY_FMT
            return (int(n) if n.is_integer() and "." not in raw else n), (INT_FMT if "." not in raw else DEC_FMT)
    return text, None


# ---------------------------------------------------------------------------------------------- بناء الملف
def safe_sheet_name(name: str, used: set[str]) -> str:
    base = _BAD_SHEET.sub(" ", clean_text(name) or "جدول").strip("' ")[:31] or "جدول"
    if base.lower() in RESERVED_SHEETS:
        base = base + "_"
    out, i = base, 2
    while out.lower() in used:
        suffix = f" ({i})"
        out = base[: 31 - len(suffix)] + suffix
        i += 1
    used.add(out.lower())
    return out


def safe_filename(name: str, default: str = "تصدير") -> str:
    base = _BAD_FILE.sub("-", clean_text(name) or default).strip(" .-")[:80] or default
    return base if base.lower().endswith(".xlsx") else base + ".xlsx"


def _put(ws, row: int, col: int, value: Any):
    """يكتب خلية؛ النص الذي يبدأ بـ = يبقى نصاً (لا يتحول لمعادلة تعرض خطأ أو تنفّذ شيئاً غير مقصود)."""
    cell = ws.cell(row=row, column=col, value=value)
    if isinstance(value, str) and value[:1] == "=":
        cell.data_type = "s"
    return cell


IMG_MAX_PX = 260          # أكبر ضلع لصورة داخل الإكسل (بكسل)


def _add_images(ws, images: Sequence[tuple[str, bytes]], col: int, start_row: int = 1) -> None:
    """يضع الصور بدءاً من (start_row, col): عنوان (اسم الصنف) ثم الصورة. حتى 6 صور عموداً واحداً، والأكثر (حتى 100) في شبكة 3 أعمدة."""
    from openpyxl.drawing.image import Image as XLImage
    from PIL import Image as PILImage
    prepared = []
    for cap, data in images:
        try:
            im = PILImage.open(io.BytesIO(data)).convert("RGB")
            im.thumbnail((IMG_MAX_PX, IMG_MAX_PX))
            buf = io.BytesIO()
            im.save(buf, "JPEG", quality=80)                  # JPEG خفيف: مئة صورة لا تضخّم الملف
            buf.seek(0)
            prepared.append((cap, XLImage(buf)))
        except Exception:  # noqa: BLE001 — صورة تالفة لا توقف التصدير
            continue
    if not prepared:
        return
    per_row = 1 if len(prepared) <= 6 else 3
    block = 2 + math.ceil(IMG_MAX_PX / 20)                  # صفوف الكتلة (عنوان + صورة) — الصف الافتراضي ≈ 20 بكسل
    for c in range(per_row):
        letter = get_column_letter(col + c)
        ws.column_dimensions[letter].width = max(ws.column_dimensions[letter].width or 0, 40)
    for i, (cap, xi) in enumerate(prepared):
        r, c = divmod(i, per_row)
        row, column = start_row + r * block, col + c
        letter = get_column_letter(column)
        title = ws.cell(row=row, column=column, value=clean_text(cap) or "صورة")
        title.font = Font(name=FONT, bold=True, size=11)
        title.alignment = Alignment(horizontal="center", wrap_text=True)
        xi.anchor = f"{letter}{row + 1}"
        ws.add_image(xi)


def build_workbook(sheets: Sequence[Sheet], business: str = "", images: Sequence[tuple[str, bytes]] | None = None) -> bytes:
    """images: [(اسم الصنف، JPEG)] تُوضع بجانب أول جدول (أو في ورقة «الصور» إن لم توجد جداول)."""
    if not sheets and not images:
        raise ValueError("لا توجد جداول للتصدير")
    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    thin = Side(style="thin", color=LINE)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_font = Font(name=FONT, bold=True, color="FFFFFF", size=11)
    head_fill = PatternFill("solid", fgColor=GREEN)
    zebra = PatternFill("solid", fgColor=MINT)
    body_font = Font(name=FONT, size=11)
    for sh in sheets:
        ws = wb.create_sheet(safe_sheet_name(sh.name, used))
        ws.sheet_view.rightToLeft = True
        ws.sheet_properties.tabColor = GREEN
        ncols = max([len(sh.headers)] + [len(r) for r in sh.rows]) if (sh.headers or sh.rows) else 1
        headers = [clean_text(h) for h in sh.headers] + [""] * (ncols - len(sh.headers))
        for c, h in enumerate(headers, 1):
            cell = _put(ws, 1, c, _fit(h))
            cell.font, cell.fill, cell.border = head_font, head_fill, border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 26
        widths = [max(8, len(h) + 4) for h in headers]
        for r, row in enumerate(sh.rows, 2):
            for c in range(ncols):
                v, fmt = coerce(row[c] if c < len(row) else None, c, sh.money_cols)
                cell = _put(ws, r, c + 1, v)
                cell.font, cell.border = body_font, border
                if fmt:
                    cell.number_format = fmt
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
                if r % 2 == 1:
                    cell.fill = zebra
                shown = (f"{v:,.2f}" if isinstance(v, float) else str(v)) if v is not None else ""
                widths[c] = min(60, max(widths[c], max((len(p) for p in shown.split("\n")), default=0) + 3))
        last = len(sh.rows) + 1
        if sh.totals and sh.rows and sh.money_cols:
            tr = last + 1
            lab = ws.cell(row=tr, column=1, value="الإجمالي")
            for c in range(1, ncols + 1):
                cell = ws.cell(row=tr, column=c)
                cell.font = Font(name=FONT, bold=True, size=11)
                cell.fill = PatternFill("solid", fgColor="D1FAE5")
                cell.border = border
            lab.alignment = Alignment(horizontal="right")
            for c in sorted(sh.money_cols):
                if c < ncols:
                    col = get_column_letter(c + 1)
                    cell = ws.cell(row=tr, column=c + 1, value=f"=SUM({col}2:{col}{last})")
                    cell.number_format = MONEY_FMT
                    cell.alignment = Alignment(horizontal="center")
        for c, wd in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(c)].width = sh.widths.get(c - 1, wd)
        if images and sh is sheets[0]:                     # الصور بجانب الجدول (عمود فارغ بعده)
            _add_images(ws, images, ncols + 2)
        ws.freeze_panes = "A2"
        if sh.rows:
            ws.auto_filter.ref = f"A1:{get_column_letter(ncols)}{last}"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        if business:
            ws.oddFooter.center.text = f"{business} — {datetime.now():%Y-%m-%d}"
    if images and not sheets:
        ws = wb.create_sheet("الصور")
        ws.sheet_view.rightToLeft = True
        _add_images(ws, images, 1)
    wb.properties.creator = business or "دفتري"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------------------------- جداول Markdown
_SEP = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")


def _cells(line: str) -> list[str]:
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|") and not s.endswith("\\|"):
        s = s[:-1]
    return [clean_text(c.replace("\\|", "|")) for c in re.split(r"(?<!\\)\|", s)]


def parse_markdown_tables(text: str) -> list[Sheet]:
    """يستخرج كل جداول Markdown من نص (ترويسة + سطر فاصل ---|--- + صفوف). اسم الورقة من العنوان الذي يسبق الجدول."""
    lines = (text or "").splitlines()
    out: list[Sheet] = []
    i, n = 0, len(lines)
    while i < n - 1:
        if "|" in lines[i] and _SEP.match(lines[i + 1]) and "-" in lines[i + 1]:
            headers = _cells(lines[i])
            title = ""
            for k in range(i - 1, max(-1, i - 4), -1):
                t = lines[k].strip()
                if t:
                    # عنوان فقط إن كان سطر عنوان (#) أو نصاً عريضاً أو ينتهي بنقطتين — لا جملة شرح عادية
                    if t.startswith("#") or t.startswith("**") or t.rstrip().endswith((":", "：")):
                        title = re.sub(r"^[#>*\-\s]+", "", t).rstrip(":：* ").strip()
                    break
            rows = []
            j = i + 2
            while j < n and "|" in lines[j] and lines[j].strip():
                rows.append(_cells(lines[j]))
                j += 1
            out.append(Sheet(title if 0 < len(title) <= 40 else f"جدول {len(out) + 1}", headers, rows))
            i = j
        else:
            i += 1
    return out


# ---------------------------------------------------------------------------------------------- مدخلات النموذج
def _cell_value(c: Any) -> Any:
    if isinstance(c, (dict, list, tuple)):
        return json.dumps(c, ensure_ascii=False)
    return c


def normalize_sheets(raw: Any) -> list[Sheet]:
    """يحوّل ما يرسله النموذج لأداة export_excel إلى أوراق صالحة مهما كان شكله.

    النماذج المختلفة (Gemini / GPT / Llama …) لا ترسل الشكل نفسه دائماً، فنقبل:
      • قائمة أوراق (الشكل المطلوب)، أو ورقة واحدة، أو {"sheets": [...]}، أو نص JSON، أو جدول Markdown نصي.
      • صفوف على شكل قوائم، أو قواميس {عنوان: قيمة} (تُستخرج منها العناوين إن غابت)، أو قيم مفردة.
    كان الصف القاموسي يُكتب سابقاً بأسماء مفاتيحه بدل قيمه دون أي تنبيه (بيانات خاطئة)."""
    if isinstance(raw, (bytes, bytearray)):
        raw = raw.decode("utf-8", "replace")
    if isinstance(raw, str):
        text = raw.strip()
        try:
            raw = json.loads(text)
        except ValueError:
            tables = parse_markdown_tables(text)
            if tables:
                return tables
            raise ValueError("sheets ليست JSON صالحاً ولا جدول Markdown — أرسل قائمة أوراق (name, headers, rows).")
    if isinstance(raw, dict):
        raw = raw["sheets"] if isinstance(raw.get("sheets"), list) else [raw]
    out: list[Sheet] = []
    for i, sh in enumerate(raw or [], 1):
        if not isinstance(sh, dict):
            continue
        rows_in = sh.get("rows") if sh.get("rows") is not None else sh.get("data")
        if isinstance(rows_in, dict):
            rows_in = [rows_in]
        rows_in = list(rows_in or [])
        headers = [str(h) for h in (sh.get("headers") or sh.get("columns") or [])]
        if not headers:                                   # عناوين من مفاتيح الصفوف القاموسية (بترتيب ظهورها)
            for r in rows_in:
                if isinstance(r, dict):
                    headers += [str(k) for k in r if str(k) not in headers]
        rows: list[list[Any]] = []
        for r in rows_in:
            if isinstance(r, dict):
                by_name = {str(k): v for k, v in r.items()}
                rows.append([_cell_value(by_name.get(h)) for h in headers])
            elif isinstance(r, (list, tuple)):
                rows.append([_cell_value(c) for c in r])
            elif r is not None:
                rows.append([_cell_value(r)])
        if not headers and not rows:
            continue
        out.append(Sheet(str(sh.get("name") or sh.get("title") or f"جدول {len(out) + 1}"), headers, rows))
    return out
